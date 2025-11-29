from modules.template_processor import load_template, save_document, fill_docx_template, extract_placeholders
from modules.llm_validator import validate_filled_docx_bytes

# Create a tiny docx
from docx import Document
from io import BytesIO

doc = Document()
doc.add_paragraph('Hello [NAME]')
buf = BytesIO()
doc.save(buf)
bytes_in = buf.getvalue()

mapping = {'NAME': 'Alice'}

# Fill and save
template = load_template(bytes_in)
filled = fill_docx_template(template, mapping)
final_bytes = save_document(filled, mapping)

# Validator without API key should return error status
res = validate_filled_docx_bytes(final_bytes, mapping, api_key=None)
print('Validator fallback:', res)
assert res['status'] == 'error'
print('Test passed: validator fallback produces error when no API key')
